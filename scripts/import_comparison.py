"""Import the supplied comparison workbook without modifying it; requires openpyxl.

Optionally verify ordering codes and prices against the two supplied catalogues.
"""
import argparse
import json
import re
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook

SOURCES = {
    'ABB': {'title': 'ABB IEC LV Motors, India Price List', 'edition': 'FRSM69A', 'effectiveDate': '2026-09-01', 'pages': [10, 11, 12, 14, 15, 16]},
    'Innomotics': {'title': 'Innomotics Motors Price List', 'edition': 'LP-210', 'effectiveDate': '2026-08-14', 'pages': [15, 16, 18, 19]},
    'CG': {'title': 'CG Motors Price List', 'edition': 'LTM27', 'effectiveDate': '2026-09-02', 'pages': [13, 14, 15], 'pageNumbering': 'Printed page number (PDF page is one higher)'}
}


def import_workbook(source, pdf_dir=None):
    sheet = load_workbook(source, data_only=True)['Source Details']
    abb = json.loads(Path('data/abb-frsm69a.json').read_text(encoding='utf-8'))['motors']
    entries = []
    for row in sheet.iter_rows(min_row=2):
        company, efficiency, poles, hp, kw, frame, mf, model, price, page, family, note = [c.value for c in row]
        manufacturer = {'ABB': 'ABB', 'INNO': 'Innomotics', 'CG': 'CG'}[company]
        available = frame != 'Not listed'
        assert hp in [10, 12.5, 15, 20, 25, 30, 40, 50] and poles in [2, 4, 6] and efficiency in ['IE3', 'IE4']
        assert (available and isinstance(price, (int, float)) and price > 0 and page) or (not available and price is None and page is None)
        documents, notes = {}, [note] if note else []
        if company == 'ABB':
            original = next(m for m in abb if m['model'] == model and m['efficiency'] == efficiency)
            assert (original['hp'], original['kw'], original['poles'], original['priceInr'], original['sourcePage']) == (hp, kw, poles, price, page), f'ABB mismatch at row {row[0].row}'
            documents = original['documents']
            notes = original['notes'] + notes
        if company == 'CG' and (str(model).endswith('*') or model == '9.30PN2_132'):
            notes.append('Catalogue * note: route this rating through Indent (special order); confirm the ordering reference with CG.')
        entries.append({'manufacturer': manufacturer, 'efficiency': efficiency, 'poles': poles,
                        'hp': hp, 'kw': kw, 'available': available, 'frame': frame if available else None,
                        'manufacturerFrame': mf if available else None, 'model': model if available else None,
                        'priceInr': int(price) if available else None, 'sourcePage': page,
                        'series': family, 'notes': notes, 'documents': documents,
                        'sourceSheet': 'Source Details', 'sourceRow': row[0].row})
    assert Counter(e['manufacturer'] for e in entries) == {'ABB': 50, 'Innomotics': 48, 'CG': 49}
    assert sum(e['available'] for e in entries) == 143
    if pdf_dir:
        from pypdf import PdfReader
        readers = {name: PdfReader(Path(pdf_dir) / filename) for name, filename in {
            'Innomotics': 'LP-210_IN Motors Price List_14th Aug 2026.pdf',
            'CG': 'LTM27 - 02.09.2026 (1).pdf'}.items()}
        texts = {}
        for e in entries:
            if e['manufacturer'] == 'ABB' or not e['available']:
                continue
            name, page = e['manufacturer'], e['sourcePage']
            if (name, page) not in texts:
                texts[name, page] = readers[name].pages[page if name == 'CG' else page-1].extract_text()
            text = texts[name, page]
            compact = re.sub(r'[\s,]', '', text)
            assert re.sub(r'\s', '', e['model']) in compact, f'Code absent from PDF: {e}'
            assert str(e['priceInr']) in compact, f'Price absent from PDF: {e}'
            if name == 'CG':
                assert re.search(re.escape(e['manufacturerFrame']) + r'\s+' + re.escape(e['model']) + r'\s+' + str(e['priceInr']), text) or re.search(re.escape(e['model']) + r'\s+' + re.escape(e['manufacturerFrame']) + r'\s+' + str(e['priceInr']), text), f'CG frame/code/price mismatch: {e}'
        print('Verified all 93 listed Innomotics/CG codes and prices against the supplied PDF pages.')
    return {'sourceWorkbook': Path(source).name, 'sources': SOURCES, 'motors': entries}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('workbook')
    parser.add_argument('--pdf-dir')
    parser.add_argument('--output', default='data/motor-catalogues.json')
    args = parser.parse_args()
    data = import_workbook(args.workbook, args.pdf_dir)
    Path(args.output).write_text(json.dumps(data, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
    print(f"Saved {len(data['motors'])} records: 143 listed motors and 4 explicit unavailable combinations.")
